#!/usr/bin/env python3
"""
Rút bảng "hàng gửi được theo tuyến" từ file training của chủ ra JSON cho web.

Nguồn: ~/Documents/MinhThien/Training NV sales.xlsx, sheet "Tổng hợp"
       (cột A mặt hàng, B ví dụ, D→M là 10 tuyến; ghi chú nằm trong comment của ô).
Đích:  src/data/hang-tuyen.json  → trang /hang-gui-duoc/<tuyến>

Chủ cập nhật file training xong thì chạy lại script này rồi deploy.

⚠️ BA GHI CHÚ CHỈ DÙNG NỘI BỘ, KHÔNG ĐƯỢC LÊN WEB (xem NOI_BO bên dưới):
   tách kiện khi hải quan soi, làm giấy tờ để chỉ bị giữ riêng mặt hàng, khuyên đi xách tay
   để né thuế. Thay bằng câu nói thẳng là có rủi ro. Thêm ghi chú mới thì soát lại danh sách này.

Dùng: python3 scripts/tao-du-lieu-hang-tuyen.py [--xlsx <file>] [--out <file>]
"""
import argparse, json, os, re, unicodedata

import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("--xlsx", default=os.path.expanduser("~/Documents/MinhThien/Training NV sales.xlsx"))
ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "..", "src", "data", "hang-tuyen.json"))
args = ap.parse_args()

# Ghi chú nội bộ → câu thay thế cho người đọc web. Khớp theo đoạn chữ đầu.
NOI_BO = [
    ("Nếu bị Hải Quan Nhật phát hiện",
     "Có rủi ro bị hải quan Nhật giữ. Minh Thiện vẫn nhận và lo giấy tờ, nhưng báo trước rủi ro cho khách."),
    ("Tách kiện",
     "Có rủi ro bị hải quan Hàn giữ. Minh Thiện vẫn nhận và tư vấn cách gửi, nhưng báo trước rủi ro cho khách."),
    ("Nên đi xách tay",
     "Hàng giá trị cao, hải quan Hàn thường yêu cầu đóng thuế. Gọi hotline để được tư vấn trước khi gửi."),
]

TRANG_THAI = {"✓": "nhan", "⚠": "dieu-kien", "✗": "khong"}


def khong_dau(s: str) -> str:
    s = s.lower().replace("đ", "d")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", s)).strip("-")


def tien(note: str) -> str:
    """Chủ gõ "300,000/kiện" — đưa lên web thì viết kiểu Việt: "300.000đ/kiện"."""
    note = re.sub(r"(\d{1,3}(?:,\d{3})+)(\s*đ)?", lambda m: m.group(1).replace(",", ".") + "đ", note)
    # "300k" cũng là 300.000đ
    return re.sub(r"\b(\d{1,3})k\b", lambda m: f"{int(m.group(1)) * 1000:,}".replace(",", ".") + "đ", note)


def loc(note: str) -> str:
    for dau, thay in NOI_BO:
        if note.startswith(dau):
            return tien(thay)
    return tien(note)


def main():
    ws = openpyxl.load_workbook(args.xlsx).worksheets[0]
    if ws.title != "Tổng hợp":
        ws = openpyxl.load_workbook(args.xlsx)["Tổng hợp"]

    tuyen = {c: str(ws.cell(3, c).value).strip() for c in range(4, 14)}
    ket_qua = {v: {"ten": v, "slug": khong_dau(v), "hang": []} for v in tuyen.values()}
    nhom = None

    for r in range(4, ws.max_row + 1):
        ten = ws.cell(r, 1).value
        if not ten:
            continue
        ten = str(ten).strip()
        # Dòng tiêu đề nhóm: chỉ có cột A
        if ws.cell(r, 4).value is None and ws.cell(r, 5).value is None:
            nhom = ten
            continue
        if ten.startswith(("Số mặt hàng", "Thời gian", "Lịch")):
            continue
        vi_du = ws.cell(r, 2).value
        for c, ten_tuyen in tuyen.items():
            o = ws.cell(r, c)
            gia_tri = str(o.value or "").strip()
            trang_thai = TRANG_THAI.get(gia_tri[:1], "chua-ro")
            ghi_chu = loc(o.comment.text.replace("\n", " ").strip()) if o.comment else ""
            ket_qua[ten_tuyen]["hang"].append({
                "ten": ten,
                "nhom": nhom,
                "viDu": str(vi_du).strip() if vi_du else "",
                "trangThai": trang_thai,
                "ghiChu": ghi_chu,
            })

    out = os.path.abspath(args.out)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(list(ket_qua.values()), f, ensure_ascii=False, indent=1)
    tong = {t: len(v["hang"]) for t, v in ket_qua.items()}
    print(f"Đã ghi {out}: {len(ket_qua)} tuyến, mỗi tuyến {set(tong.values())} mặt hàng")


if __name__ == "__main__":
    main()
